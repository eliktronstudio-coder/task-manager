"""Экспорт задач в Excel для руководителя."""
import io
from datetime import datetime
from typing import Optional
from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, and_
from sqlalchemy.orm import selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models.task import Task, TaskStatus, TaskSection

router = APIRouter(prefix="/export", tags=["export"])

STATUS_LABELS = {
    TaskStatus.IN_PROGRESS: "В работе",
    TaskStatus.CRITICAL: "Критично важно",
    TaskStatus.CLOSED: "Закрыто",
}

SECTION_LABELS = {
    TaskSection.PREPARATION: "Подготовка исходников",
    TaskSection.DELIVERY: "Поставки",
    TaskSection.QUALITY: "Качество открытия",
}


@router.get("/tasks.xlsx")
async def export_tasks_excel(
    store_id: Optional[int] = None,
    assignee_id: Optional[int] = None,
    db: AsyncSession = Depends(get_db),
):
    q = select(Task).options(
        selectinload(Task.assignee),
        selectinload(Task.creator),
        selectinload(Task.store),
    )
    filters = []
    if store_id:
        filters.append(Task.store_id == store_id)
    if assignee_id:
        filters.append(Task.assignee_id == assignee_id)
    if filters:
        q = q.where(and_(*filters))
    q = q.order_by(Task.section, Task.subsection, Task.deadline.asc().nullslast())

    result = await db.execute(q)
    tasks = result.scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Задачи"

    # Шапка
    headers = [
        "ID", "Раздел", "Подраздел", "Название", "Магазин", "Исполнитель",
        "Постановщик", "Статус", "Срок выполнения", "Создана", "Просрочена"
    ]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2B5DB8")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 30

    # Заполняем данные
    now = datetime.utcnow()
    for row, task in enumerate(tasks, 2):
        is_overdue = (
            task.deadline and task.deadline < now and task.status != TaskStatus.CLOSED
        )
        status_label = STATUS_LABELS.get(TaskStatus(task.status), task.status)
        section_label = SECTION_LABELS.get(TaskSection(task.section), task.section)

        row_data = [
            task.id,
            section_label,
            task.subsection or "",
            task.title,
            task.store.name if task.store else "",
            task.assignee.full_name if task.assignee else "Не назначен",
            task.creator.full_name if task.creator else "",
            status_label,
            task.deadline.strftime("%d.%m.%Y %H:%M") if task.deadline else "",
            task.created_at.strftime("%d.%m.%Y %H:%M"),
            "Да" if is_overdue else "Нет",
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

            # Раскрашиваем строки
            if task.status == TaskStatus.CRITICAL:
                cell.fill = PatternFill("solid", fgColor="FFCCCC")
            elif is_overdue:
                cell.fill = PatternFill("solid", fgColor="FFE5B4")
            elif task.status == TaskStatus.CLOSED:
                cell.fill = PatternFill("solid", fgColor="E8F5E9")
            elif row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F5F7FA")

    # Ширина колонок
    col_widths = [6, 25, 25, 40, 25, 25, 25, 18, 18, 18, 12]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Замораживаем первую строку
    ws.freeze_panes = "A2"

    # Добавляем лист со сводкой
    ws2 = wb.create_sheet("Сводка")
    ws2["A1"] = "Всего задач"
    ws2["B1"] = len(tasks)
    ws2["A2"] = "Закрыто"
    ws2["B2"] = sum(1 for t in tasks if t.status == TaskStatus.CLOSED)
    ws2["A3"] = "В работе"
    ws2["B3"] = sum(1 for t in tasks if t.status == TaskStatus.IN_PROGRESS)
    ws2["A4"] = "Критично важно"
    ws2["B4"] = sum(1 for t in tasks if t.status == TaskStatus.CRITICAL)
    ws2["A5"] = "Просрочено"
    ws2["B5"] = sum(1 for t in tasks if t.deadline and t.deadline < now and t.status != TaskStatus.CLOSED)
    for r in range(1, 6):
        ws2.cell(row=r, column=1).font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"tasks_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
